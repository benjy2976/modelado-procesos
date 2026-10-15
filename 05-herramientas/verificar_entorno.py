#!/usr/bin/env python3
"""Comprueba versiones, dependencias y operaciones locales con datos sinteticos."""

import json
import re
import shutil
import subprocess
import sys
from importlib.metadata import PackageNotFoundError, version
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def main() -> int:
    errors = []
    installed = {}
    if sys.version_info[:2] != (3, 13):
        errors.append("El perfil ampliado requiere Python 3.13.")
    for name, expected in re.findall(
        r"^([\w.-]+)==([^\s;\\]+)",
        (ROOT / "requirements-herramientas.txt").read_text(encoding="utf-8"),
        re.MULTILINE,
    ):
        try:
            installed[name] = version(name)
            if installed[name] != expected:
                errors.append(
                    f"{name}: instalado {installed[name]}, esperado {expected}"
                )
        except PackageNotFoundError as exc:
            errors.append(f"{name}: {exc}")
    if subprocess.run([sys.executable, "-m", "pip", "check"], check=False).returncode:
        errors.append("pip check detecto incompatibilidades.")
    try:
        import pdfplumber
        import yaml
        from docx import Document
        from jsonschema import validate
        from lxml import etree
        from openpyxl import Workbook, load_workbook
        from PIL import Image
        from pypdf import PdfReader
        from rapidfuzz.fuzz import ratio
        from reportlab.pdfgen import canvas

        data = yaml.safe_load("codigo: prueba\n")
        validate(data, {"type": "object", "required": ["codigo"]})
        etree.fromstring(
            b'<definitions xmlns="http://www.omg.org/spec/BPMN/20100524/MODEL"/>'
        )
        book = Workbook()
        book.active["A1"] = "prueba"
        stream = BytesIO()
        book.save(stream)
        stream.seek(0)
        assert load_workbook(stream).active["A1"].value == "prueba"
        document = Document()
        document.add_paragraph("prueba")
        stream = BytesIO()
        document.save(stream)
        stream.seek(0)
        assert Document(stream).paragraphs[0].text == "prueba"
        stream = BytesIO()
        pdf = canvas.Canvas(stream)
        pdf.drawString(72, 720, "Prueba del entorno")
        pdf.save()
        content = stream.getvalue()
        assert "Prueba" in PdfReader(BytesIO(content)).pages[0].extract_text()
        with pdfplumber.open(BytesIO(content)) as pdf:
            assert "Prueba" in pdf.pages[0].extract_text()
        assert Image.new("RGB", (2, 2)).size == (2, 2)
        assert ratio("prueba", "prueba") == 100
    except Exception as exc:  # noqa: BLE001 -- el diagnostico registra fallos de bibliotecas distintas
        errors.append(f"Prueba funcional: {type(exc).__name__}: {exc}")

    native = {}
    git = shutil.which("git")
    if not git and sys.platform == "win32":
        candidate = Path("C:/Program Files/Git/cmd/git.exe")
        if candidate.is_file():
            git = str(candidate)
    if git:
        for label, args in (("git", ["--version"]), ("git-lfs", ["lfs", "version"])):
            result = subprocess.run(
                [git, *args], capture_output=True, text=True, check=False
            )
            native[label] = result.stdout.strip()
            if result.returncode:
                errors.append(f"{label} no disponible.")
    else:
        errors.append("Git no encontrado. Instalarlo y abrir una terminal nueva.")

    report = {
        "python": sys.version.split()[0],
        "plataforma": sys.platform,
        "paquetes": installed,
        "herramientas": native,
        "errores": errors,
    }
    output = ROOT / "99-temporal/entorno-instalado.json"
    output.parent.mkdir(exist_ok=True)
    output.write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    for error in errors:
        print(f"ERROR: {error}")
    print(f"Paquetes comprobados: {len(installed)}; errores: {len(errors)}")
    print(f"Inventario local: {output.relative_to(ROOT)}")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
